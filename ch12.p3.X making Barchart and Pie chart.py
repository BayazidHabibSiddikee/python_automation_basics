'''Python 3.12.4 (tags/v3.12.4:8e8a4ba, Jun  6 2024, 19:30:16) [MSC v.1940 64 bit (AMD64)] on win32
Type "help", "copyright", "credits" or "license()" for more information.'''
import openpyxl
wb = openpyxl.Workbook()
sheet.active'''
Traceback (most recent call last):
  File "<pyshell#2>", line 1, in <module>
    sheet.active
NameError: name 'sheet' is not defined
>>> wb = ws.active
Traceback (most recent call last):
  File "<pyshell#3>", line 1, in <module>
    wb = ws.active
NameError: name 'ws' is not defined. Did you mean: 'wb'?'''
>>> ws = wb.active
>>> for i in range(1,11):
...     ws['a'+str(i)]=i
... 
...     
>>> ws.save('ch12.p3.x Creting Bar chart.xlsx')
Traceback (most recent call last):
  File "<pyshell#8>", line 1, in <module>
    ws.save('ch12.p3.x Creting Bar chart.xlsx')
AttributeError: 'Worksheet' object has no attribute 'save'
>>> wb.save('ch12.p3.x Creting Bar chart.xlsx')
>>> from openpyxl.chart import Barchart, PieChart, Reference, Series
Traceback (most recent call last):
  File "<pyshell#10>", line 1, in <module>
    from openpyxl.chart import Barchart, PieChart, Reference, Series
ImportError: cannot import name 'Barchart' from 'openpyxl.chart' (C:\Users\HP\AppData\Local\Programs\Python\Python312\Lib\site-packages\openpyxl\chart\__init__.py)
>>> from openpyxl.chart import BarChart, PieChart, Reference, Series
>>> ref = Reference(ws,min_col=1,min_row = 1, max_col = 1,max_row =10)
>>> serObj = Series(ref, title='First Series')
>>> wb.save('ch12.p3.x Creting Bar chart.xlsx')
>>> chart = BarChart()
>>> chart.append(serObj)
>>> wb.save('ch12.p3.x Creting Bar chart.xlsx')
>>> chart.height = 10
>>> chart.width = 15
>>> wb.save('ch12.p3.x Creting Bar chart.xlsx')
>>> chart.anchor='d1'
>>> sheet.add_chart(chart,'d1')
Traceback (most recent call last):
  File "<pyshell#22>", line 1, in <module>
    sheet.add_chart(chart,'d1')
NameError: name 'sheet' is not defined
>>> KeyboardInterrupt
>>> ws.add_chart(chart,'d1')
>>> wb.save('ch12.p3.x Creting Bar chart.xlsx')
>>> chart2 = PieChart()
>>> chart2.append(serObj)
>>> chart2.anchor='d24'
>>> ws.add_chart(chart2,'d24')
>>> wb.save('ch12.p3.x Creting Bar chart.xlsx')
