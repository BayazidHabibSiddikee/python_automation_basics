import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference, Series

wb = openpyxl.Workbook()
ws = wb.active

#Now we need some data
import random
for i in range(1,16):
    ws['a'+str(i)]=random.randint(1,100)

#Lets create reference object mind it it's obj
ref = Reference(ws, min_col=1,min_row = 1, max_col = 1, max_row = 15)
#Max row numbers or raw data be 20 only if max_row =20
#Max col means only 5 if max_col =5 else it's ok now

#Now create an series object which is aligned data to enter in charts
serObj = Series(ref,title = "Random Series")

#Lets create 2 charts
chart1 = BarChart()
chart2 = PieChart()

#Again need to enter data on different charts
chart1.append(serObj)
chart2.append(serObj)

#create the starting point
chart1.anchor = 'c1' 
chart2.anchor='c30'

#Adding chart to this object
ws.add_chart(chart1,'c1')
ws.add_chart(chart2,'c30')

#Saving the workbook to a file
wb.save('ch12.p3.x Result of barchart piechart etc.xlsx')
